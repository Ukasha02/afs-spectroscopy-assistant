"""Build the knowledge-base cache from the supplied AFS materials.

Output (gitignored, because the materials may not be redistributed):
  kb/cache/chunks.jsonl  one record per citable unit (PDF page, slide, table row, docx section)
  kb/cache/cases.json    development (D01-D08) and evaluation (H01-H06) enquiries

Run:  python -m afs_assistant.kb.build
"""

import json
import logging
import re
import sys
from dataclasses import asdict, dataclass

import docx
import openpyxl
import yaml
from docx.table import Table
from docx.text.paragraph import Paragraph
from pypdf import PdfReader

from afs_assistant import config
from afs_assistant.kb.text import normalize

PDF_SOURCES = ["S01", "S02", "S03", "S04", "S05"]


@dataclass
class Chunk:
    id: str  # stable key, e.g. "S01:p10", "TABLE:r3", "DOCX:§3"
    cite: str  # how the assistant cites it, e.g. "S01 p.10"
    source: str  # S01..S06 | TABLE | DOCX
    page: int | None
    kind: str  # page | figure_note | table_row | docx_section | link
    title: str
    text: str
    in_corpus: bool = True  # False = kept for loading cases, not put in the assessment prompt


def _load_figure_notes() -> dict[str, dict]:
    """Team-written descriptions of figures and image-only slides, keyed by 'S02:p13'."""
    raw = yaml.safe_load(config.FIGURE_NOTES_PATH.read_text(encoding="utf-8")) or {}
    return {key: value for key, value in raw.items()}


def build_pdf_chunks(source: str, figure_notes: dict[str, dict]) -> list[Chunk]:
    path = config.MATERIALS_DIR / config.SOURCE_FILES[source]
    reader = PdfReader(path)
    chunks = []
    for i, page in enumerate(reader.pages, start=1):
        key = f"{source}:p{i}"
        text = normalize(page.extract_text() or "")
        note = figure_notes.get(key)
        if note and note.get("replace_text"):
            # Extraction of this page is unusable (e.g. a scrambled table); use the note instead.
            text = f"[Page content, team transcription] {normalize(note['note'])}"
        elif note:
            # The note is appended to the page so a page citation covers both. It is labelled,
            # because it is our transcription of a figure, not the source's own wording.
            text = f"{text}\n[Figure/slide content, team transcription] {normalize(note['note'])}".strip()
        title = note["title"] if note else ""
        chunks.append(
            Chunk(id=key, cite=f"{source} p.{i}", source=source, page=i, kind="page", title=title, text=text)
        )
    return chunks


def build_table_chunks() -> list[Chunk]:
    """AFS_method_comparison_table.xlsx = machine-readable S01 Table 1 (p.6)."""
    wb = openpyxl.load_workbook(config.MATERIALS_DIR / config.SOURCE_FILES["TABLE"], read_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h) for h in rows[0]]
    chunks = []
    for n, row in enumerate(rows[1:], start=1):
        record = {h: v for h, v in zip(header, row) if v is not None}
        if not record:
            continue
        text = "; ".join(f"{k}: {v}" for k, v in record.items())
        chunks.append(
            Chunk(
                id=f"TABLE:r{n}",
                cite="S01 Table 1 (p.6)",
                source="TABLE",
                page=6,
                kind="table_row",
                title=f"{record.get('method', '')} / {record.get('measurement_technique', '')}",
                text=text,
            )
        )
    # The second sheet holds the transcription notes (units, rating meaning).
    if len(wb.worksheets) > 1:
        notes = [" ".join(str(c) for c in r if c) for r in wb.worksheets[1].iter_rows(values_only=True)]
        notes = [n for n in notes if n.strip()]
        if notes:
            chunks.append(
                Chunk(
                    id="TABLE:notes",
                    cite="S01 Table 1 (p.6)",
                    source="TABLE",
                    page=6,
                    kind="table_row",
                    title="Notes on the table transcription",
                    text=normalize(" ".join(notes)),
                )
            )
    return chunks


def _docx_blocks(document) -> list[tuple[str, str]]:
    """(style, text) for paragraphs and table rows, in document order."""
    blocks = []
    for el in document.element.body.iterchildren():
        tag = el.tag.split("}")[1]
        if tag == "p":
            p = Paragraph(el, document)
            if p.text.strip():
                blocks.append((p.style.name, p.text.strip()))
        elif tag == "tbl":
            for row in Table(el, document).rows:
                cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                blocks.append(("Table", " | ".join(cells)))
    return blocks


def build_docx_chunks() -> tuple[list[Chunk], dict[str, dict]]:
    """Split the docx by Heading 1 sections; split §1 and §2 further by D-case."""
    document = docx.Document(config.MATERIALS_DIR / config.SOURCE_FILES["DOCX"])
    blocks = _docx_blocks(document)

    sections: list[dict] = []
    current = {"h1": "Intro", "h2": None, "lines": []}
    for style, text in blocks:
        if style == "Heading 1" or (style == "Heading 2" and re.match(r"D0\d", text)):
            sections.append(current)
            h1 = text if style == "Heading 1" else current["h1"]
            current = {"h1": h1, "h2": text if style == "Heading 2" else None, "lines": []}
        else:
            current["lines"].append(text if style != "Table" else f"[table] {text}")
    sections.append(current)

    chunks, d_cases = [], {}
    for sec in sections:
        if not sec["lines"]:
            continue
        num = re.match(r"(\d)\.", sec["h1"])
        sec_no = num.group(1) if num else "0"
        case = re.match(r"(D0\d)\s*·\s*(.+)", sec["h2"] or "")
        body = normalize("\n".join(sec["lines"]))
        if case and sec_no == "1":
            # Development enquiry text: used as an input case, not as evidence.
            enquiry = sec["lines"][0]
            d_cases[case.group(1)] = {"id": case.group(1), "title": case.group(2).strip(), "enquiry": enquiry}
            chunks.append(
                Chunk(
                    id=f"DOCX:§1/{case.group(1)}",
                    cite=f"DOCX §1 {case.group(1)}",
                    source="DOCX",
                    page=None,
                    kind="docx_section",
                    title=sec["h2"],
                    text=body,
                    in_corpus=False,
                )
            )
        elif case:
            chunks.append(
                Chunk(
                    id=f"DOCX:§{sec_no}/{case.group(1)}",
                    cite=f"DOCX §{sec_no} {case.group(1)}",
                    source="DOCX",
                    page=None,
                    kind="docx_section",
                    title=f"Worked example: {sec['h2']}",
                    text=body,
                    # Answer to a different enquiry: used only as a style exemplar, never as evidence.
                    in_corpus=False,
                )
            )
        else:
            chunks.append(
                Chunk(
                    id=f"DOCX:§{sec_no}",
                    cite=f"DOCX §{sec_no}" if sec_no != "0" else "DOCX intro",
                    source="DOCX",
                    page=None,
                    kind="docx_section",
                    title=sec["h1"],
                    text=body,
                    # The challenge reference document is background (answer guidance), never evidence.
                    in_corpus=False,
                )
            )
    return chunks, d_cases


def build_h_cases() -> dict[str, dict]:
    text = (config.MATERIALS_DIR / config.SOURCE_FILES["SCENARIOS"]).read_text(encoding="utf-8")
    pattern = re.compile(r"^(H0\d) \| (.+?)\n=+\n(.*?)(?=\n=+\nH0\d|\Z)", re.S | re.M)
    cases = {}
    for case_id, title, body in pattern.findall(text):
        body = re.sub(r"\n=+\s*$", "", body.strip())
        cases[case_id] = {"id": case_id, "title": title.strip(), "enquiry": normalize(body)}
    return cases


def build_link_chunk() -> Chunk:
    url = (config.MATERIALS_DIR / config.SOURCE_FILES["S06"]).read_text(encoding="utf-8").strip()
    return Chunk(
        id="S06",
        cite="S06",
        source="S06",
        page=None,
        kind="link",
        title="Laboratory website link (business context)",
        text=f"Link to the MIR Photonics laboratory website: {url}",
    )


def main() -> int:
    logging.getLogger("pypdf").setLevel(logging.ERROR)
    missing = [f for f in config.SOURCE_FILES.values() if not (config.MATERIALS_DIR / f).exists()]
    if missing:
        print(f"Missing files in {config.MATERIALS_DIR}: {missing}\nSee materials/README.md.", file=sys.stderr)
        return 1

    figure_notes = _load_figure_notes()
    chunks: list[Chunk] = []
    for source in PDF_SOURCES:
        chunks += build_pdf_chunks(source, figure_notes)
    chunks += build_table_chunks()
    docx_chunks, d_cases = build_docx_chunks()
    chunks += docx_chunks
    chunks.append(build_link_chunk())
    h_cases = build_h_cases()

    unused = set(figure_notes) - {c.id for c in chunks}
    if unused:
        print(f"Warning: figure notes without a matching page: {sorted(unused)}", file=sys.stderr)

    config.KB_CACHE_DIR.mkdir(parents=True, exist_ok=True)
    with config.CHUNKS_PATH.open("w", encoding="utf-8") as f:
        for c in chunks:
            f.write(json.dumps(asdict(c), ensure_ascii=False) + "\n")
    (config.KB_CACHE_DIR / "cases.json").write_text(
        json.dumps({"D": d_cases, "H": h_cases}, indent=2, ensure_ascii=False), encoding="utf-8"
    )

    by_source: dict[str, int] = {}
    for c in chunks:
        by_source[c.source] = by_source.get(c.source, 0) + 1
    in_corpus = [c for c in chunks if c.in_corpus]
    print(f"Wrote {len(chunks)} chunks ({len(in_corpus)} in corpus) to {config.CHUNKS_PATH}")
    print("  per source: " + ", ".join(f"{k}={v}" for k, v in by_source.items()))
    print(f"  corpus size: {sum(len(c.text) for c in in_corpus):,} characters")
    print(f"  cases: D={sorted(d_cases)} H={sorted(h_cases)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
